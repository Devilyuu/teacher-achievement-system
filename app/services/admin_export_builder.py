import re
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import EXPORT_DIR
from app.models import AchievementStatus
from app.services.admin_summary import AdminSummary, missing_reasons
from app.services.export_builder import CATEGORY_ORDER


TEACHER_HEADERS = [
    "教师",
    "部门",
    "账号状态",
    "年度状态",
    "成果数量",
    "可申报数量",
    "待完善数量",
    "支撑材料数量",
    "申报积分合计",
]

DETAIL_HEADERS = [
    "序号",
    "教师",
    "部门",
    "项目大类",
    "项目小类",
    "成果名称",
    "申报性质",
    "申报级别",
    "本人角色",
    "申报积分",
    "状态",
    "材料数量",
]

MISSING_HEADERS = [
    "序号",
    "教师",
    "部门",
    "项目大类",
    "项目小类",
    "成果名称",
    "申报性质",
    "申报级别",
    "申报积分",
    "状态",
    "缺失原因",
]


def build_admin_summary_workbook(summary: AdminSummary) -> Path:
    output_dir = _new_export_dir(summary)
    workbook_path = output_dir / (
        f"{summary.filters.year}年度教师成果汇总_{_scope_label(summary)}.xlsx"
    )
    _write_summary_workbook(workbook_path, summary)
    return workbook_path


def build_admin_material_package(summary: AdminSummary) -> Path:
    output_dir = _new_export_dir(summary)
    workbook_path = output_dir / "00_年度教师成果汇总.xlsx"
    zip_path = output_dir / (
        f"{summary.filters.year}年度支撑材料包_{_scope_label(summary)}.zip"
    )
    _write_summary_workbook(workbook_path, summary)

    with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
        archive.write(workbook_path, workbook_path.name)
        for achievement_index, achievement in enumerate(
            summary.achievements,
            start=1,
        ):
            category_folder = _category_folder(achievement.category)
            department = _safe_name(achievement.user.department or "未设置部门")
            teacher = _safe_name(achievement.user.full_name)
            for material in achievement.materials:
                source = Path(material.stored_path)
                if not source.is_file():
                    continue
                archive_name = (
                    f"01_教师材料/{department}/{teacher}/{category_folder}/"
                    f"{_material_name(achievement_index, achievement.title, material)}"
                )
                archive.write(source, archive_name)
    return zip_path


def _write_summary_workbook(path: Path, summary: AdminSummary) -> None:
    workbook = Workbook()
    teacher_sheet = workbook.active
    teacher_sheet.title = "教师汇总"
    detail_sheet = workbook.create_sheet("成果明细")
    missing_sheet = workbook.create_sheet("材料缺失")

    teacher_sheet.append(TEACHER_HEADERS)
    for row in summary.teachers:
        teacher_sheet.append(
            [
                row.user.full_name,
                row.user.department or "未设置部门",
                "启用" if row.user.is_active else "停用",
                row.annual_state.status,
                row.achievement_count,
                row.ready_count,
                row.needs_info_count,
                row.material_count,
                row.claimed_score,
            ]
        )

    detail_sheet.append(DETAIL_HEADERS)
    missing_sheet.append(MISSING_HEADERS)
    for index, achievement in enumerate(summary.achievements, start=1):
        detail_sheet.append(
            [
                index,
                achievement.user.full_name,
                achievement.user.department or "未设置部门",
                achievement.category,
                achievement.subcategory,
                achievement.title or "未填写成果名称",
                achievement.claim_nature,
                achievement.level or "待确认",
                achievement.personal_role or "",
                achievement.claimed_score,
                achievement.status,
                len(achievement.materials),
            ]
        )
        reasons = missing_reasons(achievement)
        if reasons or achievement.status == AchievementStatus.needs_info.value:
            missing_sheet.append(
                [
                    index,
                    achievement.user.full_name,
                    achievement.user.department or "未设置部门",
                    achievement.category,
                    achievement.subcategory,
                    achievement.title or "未填写成果名称",
                    achievement.claim_nature,
                    achievement.level or "待确认",
                    achievement.claimed_score,
                    achievement.status,
                    "；".join(reasons) or "状态为待完善",
                ]
            )

    for sheet in (teacher_sheet, detail_sheet, missing_sheet):
        _format_sheet(sheet)
    workbook.save(path)


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="1F3B64")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column in range(1, sheet.max_column + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, min(sheet.max_row, 100) + 1)
        ]
        width = min(max(max(map(len, values), default=8) + 2, 10), 38)
        sheet.column_dimensions[get_column_letter(column)].width = width


def _new_export_dir(summary: AdminSummary) -> Path:
    output_dir = (
        EXPORT_DIR
        / "admin"
        / str(summary.filters.year)
        / uuid4().hex
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def _scope_label(summary: AdminSummary) -> str:
    if summary.filters.teacher_id is not None and summary.teachers:
        return _safe_name(summary.teachers[0].user.full_name)
    if summary.filters.department:
        return _safe_name(summary.filters.department)
    return "全部部门"


def _category_folder(category: str) -> str:
    category_no = CATEGORY_ORDER.index(category) + 1 if category in CATEGORY_ORDER else 99
    return f"{category_no:02d}_{_safe_name(category)}"


def _material_name(
    achievement_index: int,
    achievement_title: str,
    material,
) -> str:
    extension = material.file_ext or Path(material.original_filename).suffix
    if extension and not extension.startswith("."):
        extension = f".{extension}"
    stem = "_".join(
        [
            f"{achievement_index:03d}",
            achievement_title or "未填写成果名称",
            material.material_no,
            material.display_name,
        ]
    )
    return f"{_safe_name(stem)}{extension}"


def _safe_name(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", str(value)).strip().rstrip(".")
    return cleaned or "未命名"
