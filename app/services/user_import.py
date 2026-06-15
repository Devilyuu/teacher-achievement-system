import json
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from itsdangerous import BadSignature, URLSafeTimedSerializer
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill

from app.security import hash_password


HEADERS = ["用户名", "姓名", "部门", "角色", "初始密码"]
ROLE_ALIASES = {
    "教师": "teacher",
    "teacher": "teacher",
    "管理员": "admin",
    "admin": "admin",
}


@dataclass
class ImportRow:
    row_number: int
    username: str
    full_name: str
    department: str
    role: str
    password: str = field(repr=False)
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportResult:
    rows: list[ImportRow]

    @property
    def has_errors(self) -> bool:
        return any(row.errors for row in self.rows)


def build_user_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用户导入"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E101"
    sheet.row_dimensions[1].height = 25
    widths = [18, 16, 24, 14, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    role_validation = DataValidation(
        type="list",
        formula1='"教师,管理员"',
        allow_blank=False,
    )
    role_validation.error = "请选择教师或管理员"
    role_validation.errorTitle = "角色无效"
    sheet.add_data_validation(role_validation)
    role_validation.add("D2:D101")

    instructions = workbook.create_sheet("填写说明")
    instructions.append(["填写要求"])
    instructions.append(
        ["请勿修改“用户导入”工作表的表头；每行填写一个账号，空行会自动忽略。"]
    )
    instructions.append(
        ["20260001", "示例教师", "数字艺术学院", "教师", "teacher123"]
    )
    instructions.append(["角色可填写：教师、管理员、teacher、admin"])
    instructions.append(["初始密码至少 8 个字符；已有用户名不会被覆盖。"])
    instructions.merge_cells("A1:E1")
    instructions.merge_cells("A2:E2")
    instructions.merge_cells("A4:E4")
    instructions.merge_cells("A5:E5")
    instructions["A1"].fill = PatternFill("solid", fgColor="2563EB")
    instructions["A1"].font = Font(color="FFFFFF", bold=True)
    instructions["A1"].alignment = Alignment(horizontal="center")
    instructions["A2"].alignment = Alignment(wrap_text=True)
    for index, width in enumerate(widths, start=1):
        instructions.column_dimensions[chr(64 + index)].width = width
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse_user_import_workbook(
    content: bytes,
    existing_usernames: set[str],
) -> ImportResult:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件，请使用系统模板填写") from exc
    sheet = workbook.active
    headers = [_text(cell.value) for cell in sheet[1]]
    if headers[: len(HEADERS)] != HEADERS:
        raise ValueError("Excel 表头不正确，请重新下载系统模板")

    rows: list[ImportRow] = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True),
        start=2,
    ):
        normalized = [_text(value) for value in values]
        if not any(normalized):
            continue
        username, full_name, department, role_text, password = normalized
        role = ROLE_ALIASES.get(role_text.lower(), "")
        row = ImportRow(
            row_number=row_number,
            username=username,
            full_name=full_name,
            department=department,
            role=role,
            password=password,
        )
        _validate_required(row, role_text)
        if username in existing_usernames:
            row.errors.append("用户名已存在")
        rows.append(row)

    duplicate_names = {
        row.username
        for row in rows
        if row.username
        and sum(item.username == row.username for item in rows) > 1
    }
    for row in rows:
        if row.username in duplicate_names:
            row.errors.append("文件内用户名重复")
    return ImportResult(rows=rows)


def store_import_batch(
    rows: list[ImportRow],
    batch_dir: Path,
    secret_key: str,
) -> str:
    if any(row.errors for row in rows):
        raise ValueError("存在错误的预检结果不能生成导入批次")
    batch_dir.mkdir(parents=True, exist_ok=True)
    batch_id = uuid4().hex
    payload = [
        {
            "username": row.username,
            "full_name": row.full_name,
            "department": row.department,
            "role": row.role,
            "password_hash": hash_password(row.password),
        }
        for row in rows
    ]
    (batch_dir / f"{batch_id}.json").write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )
    return URLSafeTimedSerializer(secret_key, salt="user-import").dumps(batch_id)


def consume_import_batch(
    token: str,
    batch_dir: Path,
    secret_key: str,
) -> list[dict[str, str]]:
    try:
        batch_id = URLSafeTimedSerializer(
            secret_key,
            salt="user-import",
        ).loads(token, max_age=3600)
    except BadSignature as exc:
        raise ValueError("导入批次无效") from exc
    path = batch_dir / f"{batch_id}.json"
    if not path.is_file():
        raise ValueError("导入批次不存在")
    rows = json.loads(path.read_text(encoding="utf-8"))
    path.unlink()
    return rows


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _validate_required(row: ImportRow, role_text: str) -> None:
    fields = [
        ("用户名", row.username),
        ("姓名", row.full_name),
        ("部门", row.department),
        ("角色", role_text),
        ("初始密码", row.password),
    ]
    for label, value in fields:
        if not value:
            row.errors.append(f"{label}不能为空")
    if role_text and not row.role:
        row.errors.append("角色必须为教师或管理员")
    if row.password and len(row.password) < 8:
        row.errors.append("初始密码至少需要 8 个字符")
