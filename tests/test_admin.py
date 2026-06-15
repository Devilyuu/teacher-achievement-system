import json
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.database import SessionLocal
from app.models import PerformanceRule, Role, User
from app.security import hash_password, verify_password


def _login_admin(client: TestClient) -> None:
    client.post(
        "/login",
        data={"username": "admin", "password": "admin123456"},
        follow_redirects=False,
    )


def _create_teacher(username: str, password: str = "teacher-pass-123") -> None:
    db = SessionLocal()
    try:
        db.add(
            User(
                username=username,
                full_name="普通教师",
                department="教学部",
                role=Role.teacher.value,
                password_hash=hash_password(password),
            )
        )
        db.commit()
    finally:
        db.close()


def test_admin_users_page_requires_authentication(app):
    client = TestClient(app)

    response = client.get("/admin/users", follow_redirects=False)

    assert response.status_code in {401, 303}
    if response.status_code == 303:
        assert response.headers["location"] == "/login"


def test_non_admin_user_cannot_access_admin_users_page(app):
    username = f"teacher-{uuid4().hex}"
    _create_teacher(username)
    client = TestClient(app)
    client.post(
        "/login",
        data={"username": username, "password": "teacher-pass-123"},
        follow_redirects=False,
    )

    response = client.get("/admin/users", follow_redirects=False)

    assert response.status_code == 403


def test_admin_can_create_teacher_user_and_teacher_can_log_in(app):
    username = f"created-teacher-{uuid4().hex}"
    password = "new-teacher-pass-123"
    admin_client = TestClient(app)
    _login_admin(admin_client)

    response = admin_client.post(
        "/admin/users",
        data={
            "username": username,
            "full_name": "新建教师",
            "department": "信息工程学院",
            "role": Role.teacher.value,
            "password": password,
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith("/admin/users?")

    list_response = admin_client.get("/admin/users")
    assert list_response.status_code == 200
    assert username in list_response.text
    assert "新建教师" in list_response.text
    assert "信息工程学院" in list_response.text

    teacher_client = TestClient(app)
    login_response = teacher_client.post(
        "/login",
        data={"username": username, "password": password},
        follow_redirects=False,
    )

    assert login_response.status_code == 303
    assert login_response.headers["location"] == "/change-password"
    assert teacher_client.cookies.get("user_id")

    db = SessionLocal()
    try:
        created_user = db.query(User).filter_by(username=username).one()
        assert created_user.must_change_password is True
    finally:
        db.close()


def test_admin_can_deactivate_and_reactivate_teacher(app):
    username = f"toggle-teacher-{uuid4().hex}"
    _create_teacher(username)
    db = SessionLocal()
    try:
        teacher = db.query(User).filter_by(username=username).one()
        teacher_id = teacher.id
    finally:
        db.close()

    client = TestClient(app)
    _login_admin(client)

    deactivate = client.post(
        f"/admin/users/{teacher_id}/toggle-active",
        follow_redirects=False,
    )
    assert deactivate.status_code == 303

    blocked_login = TestClient(app).post(
        "/login",
        data={"username": username, "password": "teacher-pass-123"},
        follow_redirects=False,
    )
    assert blocked_login.status_code == 401

    reactivate = client.post(
        f"/admin/users/{teacher_id}/toggle-active",
        follow_redirects=False,
    )
    assert reactivate.status_code == 303

    allowed_login = TestClient(app).post(
        "/login",
        data={"username": username, "password": "teacher-pass-123"},
        follow_redirects=False,
    )
    assert allowed_login.status_code == 303


def test_admin_cannot_deactivate_own_account(app):
    client = TestClient(app)
    _login_admin(client)
    db = SessionLocal()
    try:
        admin = db.query(User).filter_by(username="admin").one()
        admin_id = admin.id
    finally:
        db.close()

    response = client.post(
        f"/admin/users/{admin_id}/toggle-active",
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "error=" in response.headers["location"]
    db = SessionLocal()
    try:
        assert db.get(User, admin_id).is_active is True
    finally:
        db.close()


def test_admin_can_reset_teacher_password_and_require_change(app):
    username = f"reset-teacher-{uuid4().hex}"
    _create_teacher(username)
    db = SessionLocal()
    try:
        teacher = db.query(User).filter_by(username=username).one()
        teacher_id = teacher.id
    finally:
        db.close()

    client = TestClient(app)
    _login_admin(client)
    response = client.post(
        f"/admin/users/{teacher_id}/reset-password",
        data={"new_password": "reset-pass-456"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    db = SessionLocal()
    try:
        teacher = db.get(User, teacher_id)
        assert verify_password("reset-pass-456", teacher.password_hash)
        assert teacher.must_change_password is True
    finally:
        db.close()


def test_duplicate_username_returns_page_error(app):
    client = TestClient(app)
    _login_admin(client)

    response = client.post(
        "/admin/users",
        data={
            "username": "admin",
            "full_name": "重复用户",
            "department": "测试",
            "role": Role.teacher.value,
            "password": "password123",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "error=" in response.headers["location"]


def test_admin_rules_page_includes_custom_category(app):
    client = TestClient(app)
    _login_admin(client)

    response = client.get("/admin/rules")

    assert response.status_code == 200
    assert "其他有价值工作（自定义）" in response.text


def test_admin_rules_page_only_lists_active_rules_and_shows_assignment_mode(app):
    db = SessionLocal()
    try:
        db.add(
            PerformanceRule(
                category="停用测试类别",
                subcategory="不应显示的停用规则",
                is_active=False,
                sort_order=999,
            )
        )
        db.commit()
    finally:
        db.close()

    client = TestClient(app)
    _login_admin(client)
    response = client.get("/admin/rules")

    assert response.status_code == 200
    assert "不应显示的停用规则" not in response.text
    assert "团队负责人申报并分配" in response.text
    assert "项目负责人统一赋分" in response.text


def _rule_form_data(category: str, subcategory: str) -> dict[str, str]:
    return {
        "category": category,
        "subcategory": subcategory,
        "base_rule": "基础分 2 分",
        "national_rule": "国家级 10 分",
        "provincial_rule": "省级 6 分",
        "city_rule": "市级 4 分",
        "school_rule": "校级 2 分",
        "college_rule": "学院级 1 分",
        "assignment_mode": "personal",
        "remark": "测试规则",
        "sort_order": "880",
    }


def _user_import_workbook(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用户导入"
    sheet.append(["用户名", "姓名", "部门", "角色", "初始密码"])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_admin_can_download_user_import_template(app):
    client = TestClient(app)
    _login_admin(client)

    response = client.get("/admin/users/import-template.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    workbook = load_workbook(BytesIO(response.content))
    assert [cell.value for cell in workbook["用户导入"][1]] == [
        "用户名",
        "姓名",
        "部门",
        "角色",
        "初始密码",
    ]


def test_non_admin_cannot_preview_user_import(app):
    username = f"import-teacher-{uuid4().hex}"
    _create_teacher(username)
    client = TestClient(app)
    client.post(
        "/login",
        data={"username": username, "password": "teacher-pass-123"},
        follow_redirects=False,
    )

    response = client.post(
        "/admin/users/import-preview",
        files={
            "file": (
                "users.xlsx",
                _user_import_workbook([]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 403


def test_user_import_preview_blocks_rows_with_errors(app):
    client = TestClient(app)
    _login_admin(client)

    response = client.post(
        "/admin/users/import-preview",
        files={
            "file": (
                "users.xlsx",
                _user_import_workbook(
                    [["admin", "重复管理员", "管理", "管理员", "password123"]]
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    assert "用户名已存在" in response.text
    assert 'action="/admin/users/import-confirm"' not in response.text


def test_admin_can_preview_and_confirm_user_import(app):
    suffix = uuid4().hex
    username = f"excel-teacher-{suffix}"
    client = TestClient(app)
    _login_admin(client)

    preview = client.post(
        "/admin/users/import-preview",
        files={
            "file": (
                "users.xlsx",
                _user_import_workbook(
                    [
                        [
                            username,
                            "批量导入教师",
                            "数字艺术学院",
                            "教师",
                            "teacher123",
                        ]
                    ]
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert preview.status_code == 200
    assert "预检通过" in preview.text
    token_marker = 'name="batch_token" value="'
    token_start = preview.text.index(token_marker) + len(token_marker)
    token_end = preview.text.index('"', token_start)
    token = preview.text[token_start:token_end]

    confirmed = client.post(
        "/admin/users/import-confirm",
        data={"batch_token": token},
        follow_redirects=False,
    )

    assert confirmed.status_code == 303
    assert "success=" in confirmed.headers["location"]
    db = SessionLocal()
    try:
        imported = db.query(User).filter_by(username=username).one()
        assert imported.full_name == "批量导入教师"
        assert imported.department == "数字艺术学院"
        assert imported.role == Role.teacher.value
        assert imported.must_change_password is True
        assert verify_password("teacher123", imported.password_hash)
    finally:
        db.close()


def test_non_admin_user_cannot_open_rule_create_page(app):
    username = f"rule-teacher-{uuid4().hex}"
    _create_teacher(username)
    client = TestClient(app)
    client.post(
        "/login",
        data={"username": username, "password": "teacher-pass-123"},
        follow_redirects=False,
    )

    response = client.get("/admin/rules/new", follow_redirects=False)

    assert response.status_code == 403


def test_admin_can_create_performance_rule(app):
    suffix = uuid4().hex
    category = f"新增规则类别-{suffix}"
    subcategory = f"新增规则小类-{suffix}"
    client = TestClient(app)
    _login_admin(client)

    response = client.post(
        "/admin/rules",
        data=_rule_form_data(category, subcategory),
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith("/admin/rules?")
    db = SessionLocal()
    try:
        created = (
            db.query(PerformanceRule)
            .filter_by(category=category, subcategory=subcategory)
            .one()
        )
        assert created.provincial_rule == "省级 6 分"
        assert created.is_team is False
        assert created.is_department_assigned is False
        assert created.is_active is True
    finally:
        db.query(PerformanceRule).filter_by(
            category=category,
            subcategory=subcategory,
        ).delete()
        db.commit()
        db.close()


def test_duplicate_performance_rule_returns_page_error(app):
    existing = SessionLocal()
    rule = existing.query(PerformanceRule).filter(
        PerformanceRule.is_active.is_(True)
    ).first()
    assert rule is not None
    category = rule.category
    subcategory = rule.subcategory
    existing.close()
    client = TestClient(app)
    _login_admin(client)

    response = client.post(
        "/admin/rules",
        data=_rule_form_data(category, subcategory),
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "error=" in response.headers["location"]


def test_blank_performance_rule_category_returns_page_error(app):
    client = TestClient(app)
    _login_admin(client)

    response = client.post(
        "/admin/rules",
        data=_rule_form_data("   ", "有效小类"),
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert "error=" in response.headers["location"]
    db = SessionLocal()
    try:
        assert (
            db.query(PerformanceRule)
            .filter_by(category="", subcategory="有效小类")
            .first()
            is None
        )
    finally:
        db.close()


def test_admin_can_edit_performance_rule(app):
    suffix = uuid4().hex
    category = f"编辑规则类别-{suffix}"
    subcategory = f"编辑规则小类-{suffix}"
    db = SessionLocal()
    try:
        rule = PerformanceRule(
            category=category,
            subcategory=subcategory,
            base_rule="旧规则",
            sort_order=881,
        )
        db.add(rule)
        db.commit()
        db.refresh(rule)
        rule_id = rule.id
    finally:
        db.close()
    client = TestClient(app)
    _login_admin(client)
    data = _rule_form_data(category, subcategory)
    data.update(
        {
            "base_rule": "更新后的基础规则",
            "assignment_mode": "team",
            "remark": "由团队负责人申报",
            "sort_order": "9",
        }
    )

    response = client.post(
        f"/admin/rules/{rule_id}/edit",
        data=data,
        follow_redirects=False,
    )

    assert response.status_code == 303
    db = SessionLocal()
    try:
        updated = db.get(PerformanceRule, rule_id)
        assert updated.base_rule == "更新后的基础规则"
        assert updated.is_team is True
        assert updated.is_department_assigned is False
        assert updated.remark == "由团队负责人申报"
        assert updated.sort_order == 9
    finally:
        db.query(PerformanceRule).filter_by(id=rule_id).delete()
        db.commit()
        db.close()


def test_rule_active_state_controls_teacher_form_options(app):
    suffix = uuid4().hex
    category = f"状态规则类别-{suffix}"
    subcategory = f"状态规则小类-{suffix}"
    db = SessionLocal()
    try:
        rule = PerformanceRule(
            category=category,
            subcategory=subcategory,
            is_active=True,
            sort_order=882,
        )
        db.add(rule)
        db.commit()
        db.refresh(rule)
        rule_id = rule.id
    finally:
        db.close()

    admin_client = TestClient(app)
    _login_admin(admin_client)
    disabled = admin_client.post(
        f"/admin/rules/{rule_id}/toggle-active",
        follow_redirects=False,
    )
    assert disabled.status_code == 303

    username = f"rule-option-teacher-{suffix}"
    _create_teacher(username)
    teacher_client = TestClient(app)
    teacher_client.post(
        "/login",
        data={"username": username, "password": "teacher-pass-123"},
        follow_redirects=False,
    )
    inactive_form = teacher_client.get("/achievements/new")
    serialized_subcategory = json.dumps(subcategory, ensure_ascii=True)[1:-1]
    assert serialized_subcategory not in inactive_form.text

    enabled = admin_client.post(
        f"/admin/rules/{rule_id}/toggle-active",
        follow_redirects=False,
    )
    assert enabled.status_code == 303
    active_form = teacher_client.get("/achievements/new")
    assert serialized_subcategory in active_form.text

    db = SessionLocal()
    try:
        db.query(PerformanceRule).filter_by(id=rule_id).delete()
        db.query(User).filter_by(username=username).delete()
        db.commit()
    finally:
        db.close()


def test_admin_can_filter_inactive_rules_by_keyword(app):
    suffix = uuid4().hex
    category = f"筛选规则类别-{suffix}"
    subcategory = f"筛选规则小类-{suffix}"
    db = SessionLocal()
    try:
        db.add(
            PerformanceRule(
                category=category,
                subcategory=subcategory,
                is_active=False,
                sort_order=883,
            )
        )
        db.commit()
    finally:
        db.close()
    client = TestClient(app)
    _login_admin(client)

    response = client.get(
        "/admin/rules",
        params={"status": "inactive", "q": suffix},
    )

    assert response.status_code == 200
    assert subcategory in response.text
    db = SessionLocal()
    try:
        db.query(PerformanceRule).filter_by(
            category=category,
            subcategory=subcategory,
        ).delete()
        db.commit()
    finally:
        db.close()
