from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import EXPORT_DIR
from app.database import Base, SessionLocal
from app.models import (
    Achievement,
    AchievementStatus,
    AnnualSubmission,
    ClaimNature,
    ExportRecord,
    Material,
    Role,
    User,
)
from app.security import hash_password
from app.services.annual_submission import (
    ANNUAL_STATUS_EXPORTED,
    confirm_annual_submission,
    get_annual_submission_state,
)
from app.services.export_builder import build_personal_export
from app.services.export_history import generate_personal_export


APPLICATION_HEADERS = [
    "序号",
    "项目大类",
    "项目小类",
    "申报性质",
    "项目具体名称",
    "申报级别",
    "审核认定级别",
    "本人角色",
    "赋分方式",
    "申报积分",
    "最终认定积分",
    "支撑材料编号+名称",
    "材料状态",
    "备注",
]

CATALOG_HEADERS = [
    "材料编号",
    "对应成果序号",
    "项目大类",
    "项目小类",
    "项目名称",
    "申报性质",
    "材料名称",
    "文件名",
    "文件类型",
    "上传时间",
]


def test_export_record_table_has_one_latest_record_per_user_and_year():
    table = Base.metadata.tables["export_records"]

    unique_columns = {
        tuple(column.name for column in constraint.columns)
        for constraint in table.constraints
        if constraint.__class__.__name__ == "UniqueConstraint"
    }

    assert ("user_id", "year") in unique_columns


def _create_user(db, username: str) -> User:
    user = User(
        username=username,
        full_name=f"导出教师{uuid4().hex[:6]}",
        department="Test",
        role=Role.teacher.value,
        password_hash=hash_password("password123"),
    )
    db.add(user)
    db.flush()
    return user


def _add_achievement_with_material(
    db,
    user: User,
    year: int,
    category: str,
    subcategory: str,
    source_path: Path,
    material_no: str,
) -> Achievement:
    achievement = Achievement(
        user_id=user.id,
        year=year,
        category=category,
        subcategory=subcategory,
        claim_nature=ClaimNature.result.value,
        title=f"{subcategory}成果",
        level="校级",
        personal_role="负责人",
        claimed_score=3,
        notes="测试导出",
    )
    db.add(achievement)
    db.flush()
    material = Material(
        achievement_id=achievement.id,
        material_no=material_no,
        display_name=f"{subcategory}证明",
        original_filename=source_path.name,
        stored_path=str(source_path),
        file_ext=source_path.suffix,
        file_size=source_path.stat().st_size,
        uploaded_at=datetime(2026, 5, 1, 9, 30),
    )
    achievement.materials.append(material)
    db.add(material)
    db.flush()
    return achievement


def _headers_from_zip(zip_path: Path, workbook_name: str, sheet_name: str) -> list[str]:
    with ZipFile(zip_path) as archive:
        workbook_bytes = BytesIO(archive.read(workbook_name))
    workbook = load_workbook(workbook_bytes)
    sheet = workbook[sheet_name]
    return [cell.value for cell in sheet[1]]


def test_build_personal_export_creates_zip_workbooks_and_category_materials(app, tmp_path):
    year = 2026
    standard_source = tmp_path / "standard-proof.pdf"
    custom_source = tmp_path / "custom-proof.docx"
    standard_source.write_bytes(b"%PDF-1.4 standard proof")
    custom_source.write_bytes(b"custom proof")

    db = SessionLocal()
    try:
        user = _create_user(db, f"export-builder-{uuid4().hex}")
        standard = _add_achievement_with_material(
            db,
            user,
            year,
            "教学",
            "精品在线开放课程建设（含虚拟仿真课程资源）",
            standard_source,
            "1-1",
        )
        standard.status = AchievementStatus.ready.value
        custom = _add_achievement_with_material(
            db,
            user,
            year,
            "其他有价值工作（自定义）",
            "自定义工作事项",
            custom_source,
            "2-1",
        )
        custom.status = AchievementStatus.ready.value
        incomplete = Achievement(
            user_id=user.id,
            year=year,
            category="科研与社会服务工作",
            subcategory="横向课题及项目",
            claim_nature=ClaimNature.result.value,
            title="待完善导出说明成果",
            claimed_score=0,
            status=AchievementStatus.needs_info.value,
        )
        db.add(incomplete)
        db.commit()

        zip_path = build_personal_export(db, user, year)
        user_id = user.id
        full_name = user.full_name
    finally:
        db.close()

    assert zip_path.exists()
    assert zip_path.parent == EXPORT_DIR / str(year) / str(user_id)

    with ZipFile(zip_path) as archive:
        names = archive.namelist()
        summary = archive.read("00_导出说明.txt").decode("utf-8")

    assert "00_导出说明.txt" in names
    assert "01_个人项目申报表.xlsx" in names
    assert "04_材料目录.xlsx" in names
    assert f"年度：{year}" in summary
    assert f"教师：{full_name}" in summary
    assert "成果总数：3" in summary
    assert "支撑材料总数：2" in summary
    assert "待完善成果：1" in summary
    assert "缺少材料成果：1" in summary
    assert "待完善导出说明成果" in summary
    assert "最终分值和级别认定仍以线下审核为准" in summary
    assert any(
        name.startswith("05_支撑材料/03_教学/") and name.endswith(".pdf")
        for name in names
    )
    assert any(
        name.startswith("05_支撑材料/10_其他有价值工作（自定义）/") and name.endswith(".docx")
        for name in names
    )
    assert _headers_from_zip(zip_path, "01_个人项目申报表.xlsx", "个人项目申报表") == APPLICATION_HEADERS
    assert _headers_from_zip(zip_path, "04_材料目录.xlsx", "材料目录") == CATALOG_HEADERS


def test_export_route_requires_auth_and_authenticated_user_can_download(app, tmp_path):
    client = TestClient(app)

    unauthenticated = client.get("/exports/2026/personal", follow_redirects=False)
    assert unauthenticated.status_code in {401, 303}

    source_path = tmp_path / "route-proof.pdf"
    source_path.write_bytes(b"%PDF-1.4 route proof")

    db = SessionLocal()
    try:
        user = _create_user(db, f"export-route-{uuid4().hex}")
        _add_achievement_with_material(
            db,
            user,
            2026,
            "教学",
            "精品在线开放课程建设（含虚拟仿真课程资源）",
            source_path,
            "1-1",
        )
        db.commit()
        user_id = user.id
    finally:
        db.close()

    client.cookies.set("user_id", str(user_id))
    response = client.get("/exports/2026/personal")

    assert response.status_code == 200
    assert response.headers["content-type"] in {"application/zip", "application/x-zip-compressed"}
    assert "filename" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith(".zip")
    with ZipFile(BytesIO(response.content)) as archive:
        assert "01_个人项目申报表.xlsx" in archive.namelist()

    db = SessionLocal()
    try:
        record = db.query(ExportRecord).one()
        achievement = db.query(Achievement).filter(Achievement.user_id == user_id).one()

        assert record.user_id == user_id
        assert record.year == 2026
        assert record.achievement_count == 1
        assert record.material_count == 1
        assert record.file_size == len(response.content)
        assert Path(record.file_path).exists()
        assert achievement.status == AchievementStatus.exported.value
    finally:
        db.close()


def test_personal_export_confirms_year_when_teacher_has_not_submitted(app, tmp_path):
    source_path = tmp_path / "auto-submit-proof.pdf"
    source_path.write_bytes(b"%PDF-1.4 auto submit proof")

    db = SessionLocal()
    try:
        user = _create_user(db, f"export-auto-submit-{uuid4().hex}")
        _add_achievement_with_material(
            db,
            user,
            2026,
            "教学",
            "教学成果",
            source_path,
            "1-1",
        )
        db.commit()

        generate_personal_export(db, user, 2026)
        state = get_annual_submission_state(db, user.id, 2026)
        submission_count = (
            db.query(AnnualSubmission)
            .filter_by(user_id=user.id, year=2026)
            .count()
        )

        assert submission_count == 1
        assert state.status == ANNUAL_STATUS_EXPORTED
        assert state.submitted_at is not None
        assert state.exported_at is not None
    finally:
        db.close()


def test_personal_export_marks_submitted_year_as_exported(app, tmp_path, monkeypatch):
    source_path = tmp_path / "submitted-proof.pdf"
    source_path.write_bytes(b"%PDF-1.4 submitted proof")

    from app.services import export_history

    class ExportClock:
        calls = 0

        @classmethod
        def utcnow(cls):
            cls.calls += 1
            if cls.calls == 1:
                return datetime(2026, 6, 2, 10, 0, 0)
            return datetime(2026, 6, 2, 10, 0, 5)

    monkeypatch.setattr(export_history, "datetime", ExportClock)

    db = SessionLocal()
    try:
        user = _create_user(db, f"export-submitted-{uuid4().hex}")
        _add_achievement_with_material(
            db,
            user,
            2026,
            "教学",
            "教学成果",
            source_path,
            "1-1",
        )
        db.commit()
        submission = confirm_annual_submission(db, user, 2026)
        submission.submitted_at = datetime(2026, 6, 2, 9, 0, 0)
        db.commit()
        generate_personal_export(db, user, 2026)
        user_id = user.id
        state = get_annual_submission_state(db, user_id, 2026)
        assert state.status == ANNUAL_STATUS_EXPORTED
        assert state.exported_at is not None
    finally:
        db.close()


def test_regenerating_same_year_updates_one_latest_record(app, tmp_path):
    client = TestClient(app)
    first_source = tmp_path / "first-proof.pdf"
    second_source = tmp_path / "second-proof.pdf"
    first_source.write_bytes(b"%PDF-1.4 first proof")
    second_source.write_bytes(b"%PDF-1.4 second proof")

    db = SessionLocal()
    try:
        user = _create_user(db, f"export-latest-{uuid4().hex}")
        _add_achievement_with_material(
            db,
            user,
            2026,
            "教学",
            "在线精品课程",
            first_source,
            "1-1",
        )
        db.commit()
        user_id = user.id
    finally:
        db.close()

    client.cookies.set("user_id", str(user_id))
    first_response = client.get("/exports/2026/personal")
    assert first_response.status_code == 200

    db = SessionLocal()
    try:
        first_record_id = db.query(ExportRecord).one().id
        user = db.get(User, user_id)
        _add_achievement_with_material(
            db,
            user,
            2026,
            "科研与社会服务工作",
            "横向项目",
            second_source,
            "2-1",
        )
        db.commit()
    finally:
        db.close()

    second_response = client.get("/exports/2026/personal")
    assert second_response.status_code == 200

    db = SessionLocal()
    try:
        records = db.query(ExportRecord).all()
        assert len(records) == 1
        assert records[0].id == first_record_id
        assert records[0].achievement_count == 2
        assert records[0].material_count == 2
        assert records[0].file_size == len(second_response.content)
    finally:
        db.close()


def test_export_review_page_summarizes_year_before_generation(app, tmp_path):
    client = TestClient(app)
    source_path = tmp_path / "review-proof.pdf"
    source_path.write_bytes(b"%PDF-1.4 review proof")

    db = SessionLocal()
    try:
        owner = _create_user(db, f"export-review-owner-{uuid4().hex}")
        other = _create_user(db, f"export-review-other-{uuid4().hex}")
        ready = _add_achievement_with_material(
            db,
            owner,
            2026,
            "教学",
            "精品在线开放课程建设（含虚拟仿真课程资源）",
            source_path,
            "1-1",
        )
        ready.status = AchievementStatus.ready.value
        incomplete = Achievement(
            user_id=owner.id,
            year=2026,
            category="科研与社会服务工作",
            subcategory="横向课题及项目",
            claim_nature=ClaimNature.result.value,
            title="Owner incomplete export item",
            claimed_score=0,
            status=AchievementStatus.needs_info.value,
        )
        other_incomplete = Achievement(
            user_id=other.id,
            year=2026,
            category="教学",
            subcategory="教学成果奖申报及获奖",
            claim_nature=ClaimNature.result.value,
            title="Other incomplete export item",
            claimed_score=0,
            status=AchievementStatus.needs_info.value,
        )
        db.add_all([incomplete, other_incomplete])
        db.commit()
        owner_id = owner.id
    finally:
        db.close()

    client.cookies.set("user_id", str(owner_id))
    response = client.get("/exports/2026/review")

    assert response.status_code == 200
    assert "2026 年度导出确认" in response.text
    assert "2 项成果" in response.text
    assert "1 项可导出" in response.text
    assert "1 项待完善" in response.text
    assert "1 份材料" in response.text
    assert "最终分值和级别认定仍以线下审核为准" in response.text
    assert "每个年度仅保留最新生成版本" in response.text
    assert "Owner incomplete export item" in response.text
    assert "Other incomplete export item" not in response.text
    assert "确认整理并下载" in response.text
    assert "确认生成并下载" not in response.text
    assert 'href="/achievements?year=2026"' in response.text
    assert 'href="/exports/2026/personal"' in response.text


def test_export_history_page_lists_only_owner_records_newest_first(app, tmp_path):
    client = TestClient(app)
    owner_file = tmp_path / "owner-2026.zip"
    older_file = tmp_path / "owner-2025.zip"
    other_file = tmp_path / "other-2027.zip"
    owner_file.write_bytes(b"owner 2026")
    older_file.write_bytes(b"owner 2025")
    other_file.write_bytes(b"other 2027")

    db = SessionLocal()
    try:
        owner = _create_user(db, f"export-owner-{uuid4().hex}")
        other = _create_user(db, f"export-other-{uuid4().hex}")
        db.flush()
        db.add_all(
            [
                ExportRecord(
                    user_id=owner.id,
                    year=2025,
                    file_name=older_file.name,
                    file_path=str(older_file),
                    achievement_count=2,
                    material_count=3,
                    file_size=older_file.stat().st_size,
                    generated_at=datetime(2026, 1, 2, 10, 0),
                ),
                ExportRecord(
                    user_id=owner.id,
                    year=2026,
                    file_name=owner_file.name,
                    file_path=str(owner_file),
                    achievement_count=4,
                    material_count=6,
                    file_size=owner_file.stat().st_size,
                    generated_at=datetime(2026, 6, 15, 10, 0),
                ),
                ExportRecord(
                    user_id=other.id,
                    year=2027,
                    file_name=other_file.name,
                    file_path=str(other_file),
                    achievement_count=8,
                    material_count=9,
                    file_size=other_file.stat().st_size,
                ),
            ]
        )
        db.commit()
        owner_id = owner.id
    finally:
        db.close()

    client.cookies.set("user_id", str(owner_id))
    response = client.get("/exports")

    assert response.status_code == 200
    assert "导出记录" in response.text
    assert response.text.index("2026 年度") < response.text.index("2025 年度")
    assert "2027 年度" not in response.text
    assert "4 项成果" in response.text
    assert "6 份材料" in response.text
    assert 'href="/exports/2026/review"' in response.text
    assert 'href="/exports/2025/review"' in response.text
    assert 'href="/exports/records/' in response.text


def test_export_history_page_marks_missing_files_for_regeneration(app, tmp_path):
    client = TestClient(app)
    missing_path = tmp_path / "missing.zip"

    db = SessionLocal()
    try:
        owner = _create_user(db, f"export-missing-{uuid4().hex}")
        db.flush()
        db.add(
            ExportRecord(
                user_id=owner.id,
                year=2026,
                file_name=missing_path.name,
                file_path=str(missing_path),
                achievement_count=1,
                material_count=0,
                file_size=100,
            )
        )
        db.commit()
        owner_id = owner.id
    finally:
        db.close()

    client.cookies.set("user_id", str(owner_id))
    response = client.get("/exports")

    assert response.status_code == 200
    assert "文件缺失，需重新生成" in response.text
    assert 'href="/exports/2026/review"' in response.text
    assert "/exports/records/" not in response.text


def test_recorded_export_download_requires_owner_and_existing_file(app, tmp_path):
    owner_file = tmp_path / "recorded.zip"
    owner_file.write_bytes(b"recorded export")

    db = SessionLocal()
    try:
        owner = _create_user(db, f"download-owner-{uuid4().hex}")
        other = _create_user(db, f"download-other-{uuid4().hex}")
        db.flush()
        record = ExportRecord(
            user_id=owner.id,
            year=2026,
            file_name=owner_file.name,
            file_path=str(owner_file),
            achievement_count=1,
            material_count=1,
            file_size=owner_file.stat().st_size,
        )
        db.add(record)
        db.commit()
        owner_id = owner.id
        other_id = other.id
        record_id = record.id
    finally:
        db.close()

    owner_client = TestClient(app)
    owner_client.cookies.set("user_id", str(owner_id))
    owner_response = owner_client.get(f"/exports/records/{record_id}/download")
    assert owner_response.status_code == 200
    assert owner_response.content == owner_file.read_bytes()

    other_client = TestClient(app)
    other_client.cookies.set("user_id", str(other_id))
    denied_response = other_client.get(
        f"/exports/records/{record_id}/download",
        follow_redirects=False,
    )
    assert denied_response.status_code == 404
