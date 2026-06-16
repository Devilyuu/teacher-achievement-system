from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from uuid import uuid4
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import UPLOAD_DIR
from app.database import SessionLocal
from app.models import Achievement, AchievementStatus, ClaimNature, Material, Role, User
from app.security import hash_password
from app.services.admin_export_builder import (
    build_admin_material_package,
    build_admin_summary_workbook,
)
from app.services.admin_summary import (
    SummaryFilters,
    build_admin_summary,
    missing_reasons,
)
from app.services.annual_submission import confirm_annual_submission


def _login_admin(client: TestClient) -> None:
    client.post(
        "/login",
        data={"username": "admin", "password": "admin123456"},
        follow_redirects=False,
    )


def _create_teacher(
    db,
    *,
    department: str,
    full_name: str,
    is_active: bool = True,
) -> User:
    user = User(
        username=f"summary-{uuid4().hex}",
        full_name=full_name,
        department=department,
        role=Role.teacher.value,
        password_hash=hash_password("password123"),
        is_active=is_active,
    )
    db.add(user)
    db.flush()
    return user


def _create_material_file(filename: str, content: bytes = b"proof") -> Path:
    directory = UPLOAD_DIR / "test-admin-summary" / uuid4().hex
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / filename
    path.write_bytes(content)
    return path


def _add_achievement(
    db,
    *,
    user: User,
    year: int,
    title: str,
    status: str,
    claimed_score: float = 3,
    category: str = "教学",
    subcategory: str = "教学成果奖申报及获奖",
    claim_nature: str = ClaimNature.result.value,
    current_stage: str = "",
    material_path: Path | None = None,
) -> Achievement:
    achievement = Achievement(
        user_id=user.id,
        year=year,
        category=category,
        subcategory=subcategory,
        claim_nature=claim_nature,
        title=title,
        claimed_score=claimed_score,
        current_stage=current_stage,
        status=status,
    )
    db.add(achievement)
    db.flush()

    if material_path:
        material = Material(
            achievement_id=achievement.id,
            material_no=f"{achievement.id}-1",
            display_name="获奖证明",
            original_filename=material_path.name,
            stored_path=str(material_path),
            file_ext=material_path.suffix,
            file_size=material_path.stat().st_size if material_path.exists() else 0,
        )
        achievement.materials.append(material)
        db.add(material)
        db.flush()
    return achievement


def test_summary_filters_intersect_and_aggregate_matching_records(app):
    db = SessionLocal()
    try:
        art_teacher = _create_teacher(
            db,
            department="艺术学院",
            full_name=f"艺术教师{uuid4().hex[:5]}",
        )
        other_teacher = _create_teacher(
            db,
            department="信息学院",
            full_name=f"信息教师{uuid4().hex[:5]}",
        )
        matching_material = _create_material_file("matching.pdf")
        matching = _add_achievement(
            db,
            user=art_teacher,
            year=2026,
            title="筛选命中的成果",
            status=AchievementStatus.ready.value,
            claimed_score=8,
            material_path=matching_material,
        )
        _add_achievement(
            db,
            user=art_teacher,
            year=2026,
            title="状态不匹配",
            status=AchievementStatus.needs_info.value,
        )
        _add_achievement(
            db,
            user=art_teacher,
            year=2027,
            title="年度不匹配",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("future.pdf"),
        )
        _add_achievement(
            db,
            user=other_teacher,
            year=2026,
            title="部门和教师不匹配",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("other.pdf"),
        )
        db.commit()

        result = build_admin_summary(
            db,
            SummaryFilters(
                year=2026,
                department="艺术学院",
                teacher_id=art_teacher.id,
                status=AchievementStatus.ready.value,
            ),
        )

        assert [item.id for item in result.achievements] == [matching.id]
        assert result.metrics.teacher_count == 1
        assert result.metrics.achievement_count == 1
        assert result.metrics.claimed_score == 8
        assert result.metrics.needs_info_count == 0
        assert result.metrics.material_count == 1
        assert len(result.teachers) == 1
        assert result.teachers[0].user.id == art_teacher.id
        assert result.teachers[0].ready_count == 1
    finally:
        db.close()


def test_summary_keeps_inactive_teacher_history_and_reports_missing_reasons(app):
    db = SessionLocal()
    try:
        inactive_teacher = _create_teacher(
            db,
            department="艺术学院",
            full_name=f"停用教师{uuid4().hex[:5]}",
            is_active=False,
        )
        missing_file = (
            UPLOAD_DIR
            / "test-admin-summary"
            / uuid4().hex
            / "missing-proof.pdf"
        )
        achievement = _add_achievement(
            db,
            user=inactive_teacher,
            year=2026,
            title="",
            status=AchievementStatus.needs_info.value,
            claimed_score=0,
            claim_nature=ClaimNature.process.value,
            current_stage="",
            material_path=missing_file,
        )
        db.commit()

        result = build_admin_summary(
            db,
            SummaryFilters(year=2026, department="艺术学院"),
        )
        reasons = missing_reasons(achievement)

        assert achievement.id in [item.id for item in result.achievements]
        teacher_row = next(
            row for row in result.teachers if row.user.id == inactive_teacher.id
        )
        assert teacher_row.user.is_active is False
        assert teacher_row.needs_info_count == 1
        assert "申报信息不完整" in reasons
        assert "申报积分未填写或不大于零" in reasons
        assert "过程性工作缺少进展说明" in reasons
        assert "材料文件不存在" in reasons
        assert "未上传支撑材料" not in reasons
    finally:
        db.close()


def test_summary_excludes_achievements_owned_by_admin_accounts(app):
    db = SessionLocal()
    try:
        admin = db.query(User).filter_by(username="admin").one()
        admin_achievement = _add_achievement(
            db,
            user=admin,
            year=2026,
            title=f"管理员测试成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("admin-proof.pdf"),
        )
        db.commit()

        result = build_admin_summary(db, SummaryFilters(year=2026))

        assert admin_achievement.id not in [
            achievement.id for achievement in result.achievements
        ]
    finally:
        db.close()


def test_admin_summary_page_requires_admin(app):
    unauthenticated = TestClient(app).get(
        "/admin/summary",
        follow_redirects=False,
    )
    assert unauthenticated.status_code in {303, 401}

    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department="权限测试学院",
            full_name=f"普通教师{uuid4().hex[:5]}",
        )
        db.commit()
        teacher_id = teacher.id
    finally:
        db.close()

    teacher_client = TestClient(app)
    teacher_client.cookies.set("user_id", str(teacher_id))
    forbidden = teacher_client.get("/admin/summary")
    assert forbidden.status_code == 403


def test_admin_summary_page_renders_filtered_metrics_rows_and_export_links(app):
    client = TestClient(app)
    _login_admin(client)
    department = f"汇总学院{uuid4().hex[:5]}"

    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department=department,
            full_name=f"汇总教师{uuid4().hex[:5]}",
        )
        matching = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"页面命中成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            claimed_score=6,
            material_path=_create_material_file("page-proof.pdf"),
        )
        excluded = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"页面排除成果{uuid4().hex[:5]}",
            status=AchievementStatus.needs_info.value,
        )
        db.commit()
        teacher_id = teacher.id
        teacher_name = teacher.full_name
        matching_title = matching.title
        excluded_title = excluded.title
    finally:
        db.close()

    response = client.get(
        "/admin/summary",
        params={
            "year": 2026,
            "department": department,
            "teacher_id": teacher_id,
            "status": AchievementStatus.ready.value,
        },
    )

    assert response.status_code == 200
    assert "年度汇总" in response.text
    assert teacher_name in response.text
    assert matching_title in response.text
    assert excluded_title not in response.text
    assert 'data-metric="teacher-count">1<' in response.text
    assert 'data-metric="achievement-count">1<' in response.text
    assert 'data-metric="claimed-score">6.0<' in response.text
    assert "/admin/summary/export.xlsx?" in response.text
    assert "/admin/summary/materials.zip?" in response.text
    assert "year=2026" in response.text
    assert f"teacher_id={teacher_id}" in response.text
    assert f"department={quote(department)}" in response.text
    assert f"status={quote(AchievementStatus.ready.value)}" in response.text


def test_admin_summary_page_shows_annual_status_filter_and_teacher_status(app):
    client = TestClient(app)
    _login_admin(client)

    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department="年度状态学院",
            full_name=f"年度状态教师{uuid4().hex[:5]}",
        )
        _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"整理中成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
        )
        db.commit()
    finally:
        db.close()

    response = client.get("/admin/summary", params={"year": 2026})

    assert response.status_code == 200
    assert "年度状态" in response.text
    assert "全部年度状态" in response.text
    assert "整理中" in response.text


def test_admin_summary_filters_not_started_teachers_by_annual_status(app):
    db = SessionLocal()
    try:
        not_started = _create_teacher(
            db,
            department="年度筛选学院",
            full_name=f"未开始教师{uuid4().hex[:5]}",
        )
        in_progress = _create_teacher(
            db,
            department="年度筛选学院",
            full_name=f"整理中教师{uuid4().hex[:5]}",
        )
        _add_achievement(
            db,
            user=in_progress,
            year=2026,
            title=f"已有成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
        )
        db.commit()

        result = build_admin_summary(
            db,
            SummaryFilters(
                year=2026,
                department="年度筛选学院",
                annual_status="未开始",
            ),
        )

        assert [row.user.id for row in result.teachers] == [not_started.id]
        assert result.achievements == []
    finally:
        db.close()


def test_admin_summary_filters_submitted_teachers_and_their_achievements(app):
    db = SessionLocal()
    try:
        submitted = _create_teacher(
            db,
            department="已提交学院",
            full_name=f"已提交教师{uuid4().hex[:5]}",
        )
        submitted_achievement = _add_achievement(
            db,
            user=submitted,
            year=2026,
            title=f"已提交成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
        )
        in_progress = _create_teacher(
            db,
            department="已提交学院",
            full_name=f"整理中教师{uuid4().hex[:5]}",
        )
        _add_achievement(
            db,
            user=in_progress,
            year=2026,
            title=f"整理中成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
        )
        db.commit()
        confirm_annual_submission(db, submitted, 2026)

        result = build_admin_summary(
            db,
            SummaryFilters(
                year=2026,
                department="已提交学院",
                annual_status="已提交",
            ),
        )

        assert [row.user.id for row in result.teachers] == [submitted.id]
        assert [achievement.id for achievement in result.achievements] == [
            submitted_achievement.id
        ]
        assert result.metrics.teacher_count == 1
        assert result.metrics.submitted_teacher_count == 1
    finally:
        db.close()


def test_admin_can_open_teacher_achievement_as_read_only(app):
    client = TestClient(app)
    _login_admin(client)

    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department="只读学院",
            full_name=f"只读教师{uuid4().hex[:5]}",
        )
        achievement = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"只读成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("readonly.pdf"),
        )
        db.commit()
        achievement_id = achievement.id
        title = achievement.title
        teacher_name = teacher.full_name
    finally:
        db.close()

    response = client.get(f"/admin/achievements/{achievement_id}")

    assert response.status_code == 200
    assert title in response.text
    assert teacher_name in response.text
    assert f"/achievements/{achievement_id}/edit" not in response.text
    assert 'action="/materials/upload"' not in response.text
    assert f'action="/achievements/{achievement_id}/delete"' not in response.text


def test_admin_can_preview_teacher_material_inline(app):
    client = TestClient(app)
    _login_admin(client)
    material_path = _create_material_file("admin-preview.pdf", b"%PDF-1.4 admin")

    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department="Preview",
            full_name=f"Preview Teacher {uuid4().hex[:5]}",
        )
        achievement = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"Preview achievement {uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            material_path=material_path,
        )
        db.commit()
        material_id = achievement.materials[0].id
    finally:
        db.close()

    response = client.get(f"/admin/materials/{material_id}/preview")

    assert response.status_code == 200
    assert response.content == material_path.read_bytes()
    assert response.headers["content-type"].startswith("application/pdf")
    assert "inline" in response.headers["content-disposition"]


def test_admin_achievement_detail_shows_preview_for_supported_materials(app):
    client = TestClient(app)
    _login_admin(client)
    preview_path = _create_material_file("admin-detail-preview.jpg", b"jpg")
    unsupported_path = _create_material_file("admin-detail-sheet.xlsx", b"xlsx")

    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department="Preview Detail",
            full_name=f"Preview Detail Teacher {uuid4().hex[:5]}",
        )
        achievement = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"Preview detail achievement {uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            material_path=preview_path,
        )
        material = Material(
            achievement_id=achievement.id,
            material_no=f"{achievement.id}-2",
            display_name="Excel proof",
            original_filename=unsupported_path.name,
            stored_path=str(unsupported_path),
            file_ext=unsupported_path.suffix,
            file_size=unsupported_path.stat().st_size,
        )
        achievement.materials.append(material)
        db.add(material)
        db.commit()
        achievement_id = achievement.id
        preview_id = achievement.materials[0].id
        unsupported_id = material.id
    finally:
        db.close()

    response = client.get(f"/admin/achievements/{achievement_id}")

    assert response.status_code == 200
    assert f"/admin/materials/{preview_id}/preview" in response.text
    assert f"/admin/materials/{unsupported_id}/preview" not in response.text


def test_admin_summary_workbook_contains_three_filtered_sheets(app):
    department = f"导出学院{uuid4().hex[:5]}"
    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department=department,
            full_name=f"导出教师{uuid4().hex[:5]}",
        )
        ready = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"完整成果{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            claimed_score=7,
            material_path=_create_material_file("complete.pdf"),
        )
        incomplete = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"缺材料成果{uuid4().hex[:5]}",
            status=AchievementStatus.needs_info.value,
            claimed_score=2,
        )
        db.commit()
        confirm_annual_submission(db, teacher, 2026)
        summary = build_admin_summary(
            db,
            SummaryFilters(year=2026, department=department),
        )

        workbook_path = build_admin_summary_workbook(summary)
        workbook = load_workbook(workbook_path)

        assert workbook.sheetnames == ["教师汇总", "成果明细", "材料缺失"]
        teacher_sheet = workbook["教师汇总"]
        detail_sheet = workbook["成果明细"]
        missing_sheet = workbook["材料缺失"]
        assert teacher_sheet.max_row == 2
        assert teacher_sheet.cell(2, 1).value == teacher.full_name
        assert teacher_sheet.cell(1, 4).value == "年度状态"
        assert teacher_sheet.cell(2, 4).value == "已提交"
        assert teacher_sheet.cell(2, 5).value == 2
        assert teacher_sheet.cell(2, 9).value == 9
        detail_titles = [detail_sheet.cell(row, 6).value for row in range(2, 4)]
        assert ready.title in detail_titles
        assert incomplete.title in detail_titles
        assert missing_sheet.max_row == 2
        assert missing_sheet.cell(2, 6).value == incomplete.title
        assert "未上传支撑材料" in missing_sheet.cell(2, 11).value
        assert str(2026) in workbook_path.name
        assert department in workbook_path.name
    finally:
        db.close()


def test_empty_admin_summary_workbook_still_contains_headers(app):
    db = SessionLocal()
    try:
        summary = build_admin_summary(db, SummaryFilters(year=2099))
        workbook = load_workbook(build_admin_summary_workbook(summary))

        assert workbook["教师汇总"].max_row == 1
        teacher_headers = [
            workbook["教师汇总"].cell(1, column).value
            for column in range(1, workbook["教师汇总"].max_column + 1)
        ]
        assert "年度状态" in teacher_headers
        assert workbook["成果明细"].max_row == 1
        assert workbook["材料缺失"].max_row == 1
    finally:
        db.close()


def test_admin_material_package_groups_files_and_reports_missing_physical_file(app):
    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department="艺术/学院",
            full_name=f"张:老师{uuid4().hex[:4]}",
        )
        existing = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title="获奖*项目",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("award.pdf"),
        )
        missing_path = (
            UPLOAD_DIR
            / "test-admin-summary"
            / uuid4().hex
            / "lost-proof.pdf"
        )
        missing = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title="文件丢失项目",
            status=AchievementStatus.needs_info.value,
            material_path=missing_path,
        )
        db.commit()
        summary = build_admin_summary(
            db,
            SummaryFilters(
                year=2026,
                department=teacher.department,
                teacher_id=teacher.id,
            ),
        )

        package_path = build_admin_material_package(summary)

        with ZipFile(package_path) as archive:
            names = archive.namelist()
            workbook = load_workbook(
                BytesIO(archive.read("00_年度教师成果汇总.xlsx"))
            )

        material_prefix = (
            f"01_教师材料/艺术_学院/"
            f"{teacher.full_name.replace(':', '_')}/03_教学/"
        )
        matching_names = [
            name
            for name in names
            if name.startswith(material_prefix) and name.endswith(".pdf")
        ]
        assert len(matching_names) == 1
        assert "获奖_项目" in matching_names[0]
        assert existing.materials[0].material_no in matching_names[0]
        assert not any("文件丢失项目" in name for name in names)
        missing_sheet = workbook["材料缺失"]
        teacher_sheet = workbook["教师汇总"]
        teacher_headers = [
            teacher_sheet.cell(1, column).value
            for column in range(1, teacher_sheet.max_column + 1)
        ]
        assert "年度状态" in teacher_headers
        missing_rows = [
            [cell.value for cell in row]
            for row in missing_sheet.iter_rows(min_row=2)
        ]
        assert any(
            row[5] == missing.title and "材料文件不存在" in row[10]
            for row in missing_rows
        )
    finally:
        db.close()


def test_admin_summary_export_routes_apply_filters_and_require_admin(app):
    department = f"路由导出学院{uuid4().hex[:5]}"
    db = SessionLocal()
    try:
        teacher = _create_teacher(
            db,
            department=department,
            full_name=f"路由导出教师{uuid4().hex[:5]}",
        )
        included = _add_achievement(
            db,
            user=teacher,
            year=2026,
            title=f"导出命中{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("route-export.pdf"),
        )
        excluded = _add_achievement(
            db,
            user=teacher,
            year=2027,
            title=f"导出排除{uuid4().hex[:5]}",
            status=AchievementStatus.ready.value,
            material_path=_create_material_file("route-future.pdf"),
        )
        db.commit()
        teacher_id = teacher.id
        included_title = included.title
        excluded_title = excluded.title
    finally:
        db.close()

    teacher_client = TestClient(app)
    teacher_client.cookies.set("user_id", str(teacher_id))
    assert teacher_client.get("/admin/summary/export.xlsx").status_code == 403
    assert teacher_client.get("/admin/summary/materials.zip").status_code == 403

    admin_client = TestClient(app)
    _login_admin(admin_client)
    params = {
        "year": 2026,
        "department": department,
        "teacher_id": teacher_id,
        "status": AchievementStatus.ready.value,
    }
    workbook_response = admin_client.get(
        "/admin/summary/export.xlsx",
        params=params,
    )
    package_response = admin_client.get(
        "/admin/summary/materials.zip",
        params=params,
    )

    assert workbook_response.status_code == 200
    workbook = load_workbook(BytesIO(workbook_response.content))
    detail_titles = [
        workbook["成果明细"].cell(row, 6).value
        for row in range(2, workbook["成果明细"].max_row + 1)
    ]
    assert detail_titles == [included_title]
    assert excluded_title not in detail_titles

    assert package_response.status_code == 200
    with ZipFile(BytesIO(package_response.content)) as archive:
        assert "00_年度教师成果汇总.xlsx" in archive.namelist()
        assert any(name.endswith(".pdf") for name in archive.namelist())
