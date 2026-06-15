import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.services.user_import import (
    build_user_import_template,
    consume_import_batch,
    parse_user_import_workbook,
    store_import_batch,
)


HEADERS = ["用户名", "姓名", "部门", "角色", "初始密码"]


def _workbook_bytes(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用户导入"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_user_import_template_has_expected_headers_and_example():
    content = build_user_import_template()
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["用户导入"]

    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet["A2"].value is None
    assert sheet.freeze_panes == "A2"
    instructions = workbook["填写说明"]
    assert instructions["A3"].value == "20260001"
    assert instructions["D3"].value == "教师"
    assert list(sheet.data_validations.dataValidation)


def test_parser_normalizes_valid_rows_and_skips_empty_rows():
    result = parse_user_import_workbook(
        _workbook_bytes(
            [
                ["20260001", "张老师", "数字艺术学院", "教师", "teacher123"],
                ["", "", "", "", ""],
                ["admin2", "李管理员", "数字艺术学院", "admin", "adminpass"],
            ]
        ),
        existing_usernames=set(),
    )

    assert result.has_errors is False
    assert [row.username for row in result.rows] == ["20260001", "admin2"]
    assert [row.role for row in result.rows] == ["teacher", "admin"]


def test_parser_reports_row_level_validation_errors():
    result = parse_user_import_workbook(
        _workbook_bytes(
            [
                ["existing", "已有教师", "学院", "教师", "teacher123"],
                ["duplicate", "教师甲", "学院", "教师", "teacher123"],
                ["duplicate", "教师乙", "学院", "教师", "teacher123"],
                ["short", "短密码", "学院", "教师", "123"],
                ["bad-role", "错误角色", "学院", "院长", "teacher123"],
                ["missing", "", "学院", "教师", "teacher123"],
            ]
        ),
        existing_usernames={"existing"},
    )

    assert result.has_errors is True
    errors = {row.row_number: row.errors for row in result.rows}
    assert "用户名已存在" in errors[2]
    assert "文件内用户名重复" in errors[3]
    assert "文件内用户名重复" in errors[4]
    assert "初始密码至少需要 8 个字符" in errors[5]
    assert "角色必须为教师或管理员" in errors[6]
    assert "姓名不能为空" in errors[7]


def test_import_batch_stores_hashes_and_is_consumed_once(tmp_path):
    result = parse_user_import_workbook(
        _workbook_bytes(
            [["20260001", "张老师", "数字艺术学院", "教师", "teacher123"]]
        ),
        existing_usernames=set(),
    )

    token = store_import_batch(result.rows, tmp_path, "test-secret")
    batch_files = list(tmp_path.glob("*.json"))
    assert len(batch_files) == 1
    stored_text = batch_files[0].read_text(encoding="utf-8")
    stored = json.loads(stored_text)
    assert "teacher123" not in stored_text
    assert stored[0]["password_hash"]

    rows = consume_import_batch(token, tmp_path, "test-secret")
    assert rows[0]["username"] == "20260001"
    assert list(tmp_path.glob("*.json")) == []

    with pytest.raises(ValueError, match="导入批次不存在"):
        consume_import_batch(token, tmp_path, "test-secret")
