from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import Achievement, AchievementStatus, ExportRecord, Role, User


def _login_admin(client: TestClient) -> None:
    response = client.post(
        "/login",
        data={"username": "admin", "password": "admin123456"},
        follow_redirects=False,
    )
    assert response.status_code == 303


def test_college_trial_flow_from_account_to_admin_exports(app):
    year = 2026
    suffix = uuid4().hex[:8]
    username = f"trial-teacher-{suffix}"
    initial_password = "trial-pass-123"
    new_password = "trial-new-pass-123"
    department = f"Trial College {suffix}"
    teacher_name = f"Trial Teacher {suffix}"
    achievement_title = f"Digital Art Award {suffix}"

    admin_client = TestClient(app)
    _login_admin(admin_client)

    created = admin_client.post(
        "/admin/users",
        data={
            "username": username,
            "full_name": teacher_name,
            "department": department,
            "role": Role.teacher.value,
            "password": initial_password,
        },
        follow_redirects=False,
    )
    assert created.status_code == 303

    teacher_client = TestClient(app)
    login = teacher_client.post(
        "/login",
        data={"username": username, "password": initial_password},
        follow_redirects=False,
    )
    assert login.status_code == 303
    assert login.headers["location"] == "/change-password"

    password_changed = teacher_client.post(
        "/change-password",
        data={
            "current_password": initial_password,
            "new_password": new_password,
            "confirm_password": new_password,
        },
        follow_redirects=False,
    )
    assert password_changed.status_code == 303
    assert password_changed.headers["location"] == "/"

    created_achievement = teacher_client.post(
        "/achievements",
        data={
            "year": str(year),
            "category": "Teaching",
            "subcategory": "Skills Competition Award",
            "claim_nature": "Result",
            "title": achievement_title,
            "date_range": "2026-05",
            "level": "Provincial",
            "personal_role": "Team member",
            "current_stage": "Award certificate received",
            "base_score": "0",
            "performance_score": "0",
            "claimed_score": "5",
            "notes": "Offline review confirms the final score.",
        },
        follow_redirects=False,
    )
    assert created_achievement.status_code == 303
    achievement_location = created_achievement.headers["location"]
    assert achievement_location.startswith("/achievements/")
    achievement_id = int(achievement_location.rsplit("/", 1)[1])

    uploaded = teacher_client.post(
        "/materials/upload",
        data={
            "achievement_id": str(achievement_id),
            "display_name": "Award certificate",
            "description": "Trial proof",
        },
        files={
            "file": (
                "award-certificate.pdf",
                b"%PDF-1.4 trial proof",
                "application/pdf",
            )
        },
        follow_redirects=False,
    )
    assert uploaded.status_code == 303
    assert uploaded.headers["location"].startswith(
        f"/achievements/{achievement_id}?"
    )

    export_review = teacher_client.get(f"/exports/{year}/review")
    assert export_review.status_code == 200
    assert achievement_title in export_review.text

    personal_export = teacher_client.get(f"/exports/{year}/personal")
    assert personal_export.status_code == 200
    assert personal_export.headers["content-type"] in {
        "application/zip",
        "application/x-zip-compressed",
    }
    with ZipFile(BytesIO(personal_export.content)) as archive:
        export_names = archive.namelist()
    assert any(name.endswith(".pdf") for name in export_names)
    assert any(name.endswith(".xlsx") for name in export_names)
    assert any(name.endswith(".txt") for name in export_names)

    db = SessionLocal()
    try:
        teacher = db.query(User).filter_by(username=username).one()
        achievement = db.get(Achievement, achievement_id)
        export_record = (
            db.query(ExportRecord)
            .filter_by(user_id=teacher.id, year=year)
            .one()
        )
        assert teacher.must_change_password is False
        assert achievement.user_id == teacher.id
        assert achievement.status == AchievementStatus.exported.value
        assert len(achievement.materials) == 1
        assert export_record.achievement_count == 1
        assert export_record.material_count == 1
    finally:
        db.close()

    summary_response = admin_client.get(
        "/admin/summary/export.xlsx",
        params={"year": year, "department": department},
    )
    assert summary_response.status_code == 200
    summary_workbook = load_workbook(BytesIO(summary_response.content))
    assert teacher_name in [
        summary_workbook["教师汇总"].cell(row, 1).value
        for row in range(2, summary_workbook["教师汇总"].max_row + 1)
    ]
    assert achievement_title in [
        summary_workbook["成果明细"].cell(row, 6).value
        for row in range(2, summary_workbook["成果明细"].max_row + 1)
    ]

    package_response = admin_client.get(
        "/admin/summary/materials.zip",
        params={"year": year, "department": department},
    )
    assert package_response.status_code == 200
    with ZipFile(BytesIO(package_response.content)) as archive:
        package_names = archive.namelist()
    assert any(name.endswith(".xlsx") for name in package_names)
    assert any(name.endswith(".pdf") for name in package_names)
